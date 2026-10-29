"""Выгрузка вакансий и компаний в CSV и Excel.

Лист описывается одной структурой Sheet: набор колонок, функция строки и
подготовка queryset. Поэтому новая выгрузка - это ещё один Sheet, а не копия
всей машинерии записи файла.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime
from typing import Callable

from django.db.models import Count
from django.http import HttpResponse
from django.utils import timezone

VACANCY_COLUMNS: list[tuple[str, str]] = [
    ("job_id", "LinkedIn job ID"),
    ("title", "Вакансия"),
    ("company_display", "Компания"),
    ("location", "Локация"),
    ("posted_raw", "Размещена (текст)"),
    ("posted_at", "Дата размещения"),
    ("is_reposted", "Повторно"),
    ("applicants_raw", "Откликов"),
    ("employment_type", "Тип занятости"),
    ("seniority", "Уровень"),
    # "Функция" (job function) на гостевых страницах LinkedIn не отдаётся вообще -
    # в выгрузке была бы всегда пустой колонкой. Поле в модели осталось: если
    # LinkedIn когда-нибудь начнёт его отдавать, оно заполнится и будет видно в админке.
    ("industries", "Отрасли вакансии"),
    ("benefits", "Плюшки"),
    ("status_display", "Статус обработки"),
    ("detail_error", "Ошибка обработки"),
    ("url", "Ссылка на вакансию"),
    ("company_url", "Ссылка на компанию"),
    ("company_industry", "Отрасль компании"),
    ("company_size", "Размер компании"),
    ("company_headquarters", "Штаб-квартира"),
    ("company_website", "Сайт компании"),
    ("company_followers", "Подписчиков"),
    ("company_tagline", "Компания: короткое описание"),
    ("company_description", "Компания: описание"),
    ("queries", "Найдена по запросам"),
    ("times_seen", "Раз встречена"),
    ("created_at", "Добавлена"),
    ("description", "Описание вакансии"),
]

# Старое имя: на него могли ссылаться снаружи.
COLUMNS = VACANCY_COLUMNS

COMPANY_COLUMNS: list[tuple[str, str]] = [
    ("name", "Компания"),
    ("slug", "LinkedIn slug"),
    ("linkedin_url", "Ссылка на компанию"),
    ("website", "Сайт"),
    ("industry", "Отрасль"),
    ("company_size", "Размер"),
    ("headquarters", "Штаб-квартира"),
    ("followers", "Подписчиков"),
    ("vacancies_total", "Вакансий в базе"),
    ("profile_state", "Профиль"),
    ("fetch_error", "Ошибка загрузки профиля"),
    ("details_fetched_at", "Профиль загружен"),
    ("created_at", "Добавлена"),
    ("logo_url", "Логотип"),
    ("tagline", "Короткое описание"),
    ("description", "Описание компании"),
]


def _naive(value):
    """Excel не умеет timezone-aware datetime."""
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def vacancy_row(vacancy) -> dict:
    company = vacancy.company
    return {
        "job_id": vacancy.job_id,
        "title": vacancy.title,
        "company_display": (company.name if company else "") or vacancy.company_name_raw,
        "location": vacancy.location,
        "posted_raw": vacancy.posted_raw,
        "posted_at": vacancy.posted_at,
        "is_reposted": "да" if vacancy.is_reposted else "нет",
        "applicants_raw": vacancy.applicants_raw,
        "employment_type": vacancy.employment_type,
        "seniority": vacancy.seniority,
        "job_function": vacancy.job_function,
        "industries": vacancy.industries,
        "benefits": vacancy.benefits,
        "status_display": vacancy.get_status_display(),
        "detail_error": vacancy.detail_error,
        "url": vacancy.url,
        "company_url": (company.linkedin_url if company else ""),
        "company_industry": (company.industry if company else ""),
        "company_size": (company.company_size if company else ""),
        "company_headquarters": (company.headquarters if company else ""),
        "company_website": (company.website if company else ""),
        "company_followers": (company.followers if company else ""),
        "company_tagline": (company.tagline if company else ""),
        "company_description": (company.description if company else ""),
        "queries": ", ".join(q.title for q in vacancy.search_queries.all()),
        "times_seen": vacancy.times_seen,
        "created_at": _naive(vacancy.created_at),
        "description": vacancy.description,
    }


def company_row(company) -> dict:
    if company.fetch_error:
        state = "Ошибка"
    elif company.details_fetched_at:
        state = "Загружен"
    else:
        state = "Не загружен"
    return {
        "name": company.name,
        "slug": company.slug,
        "linkedin_url": company.linkedin_url or f"https://www.linkedin.com/company/{company.slug}/",
        "website": company.website,
        "industry": company.industry,
        "company_size": company.company_size,
        "headquarters": company.headquarters,
        "followers": company.followers,
        "vacancies_total": getattr(company, "vacancy_total", None),
        "profile_state": state,
        "fetch_error": company.fetch_error,
        "details_fetched_at": _naive(company.details_fetched_at),
        "created_at": _naive(company.created_at),
        "logo_url": company.logo_url,
        "tagline": company.tagline,
        "description": company.description,
    }


def _prepare_vacancies(queryset):
    return queryset.select_related("company").prefetch_related("search_queries")


def _prepare_companies(queryset):
    # Админка уже вешает этот аннотейт в get_queryset, а команда из консоли - нет.
    if "vacancy_total" in queryset.query.annotations:
        return queryset
    return queryset.annotate(vacancy_total=Count("vacancies"))


@dataclass(frozen=True)
class Sheet:
    title: str
    file_prefix: str
    columns: list[tuple[str, str]]
    row: Callable[[object], dict]
    prepare: Callable[[object], object]
    widths: dict[str, int]


VACANCIES = Sheet(
    title="Вакансии",
    file_prefix="vacancies",
    columns=VACANCY_COLUMNS,
    row=vacancy_row,
    prepare=_prepare_vacancies,
    widths={
        "Вакансия": 45,
        "Компания": 28,
        "Локация": 28,
        "Ссылка на вакансию": 50,
        "Ссылка на компанию": 40,
        "Компания: описание": 60,
        "Описание вакансии": 80,
    },
)

COMPANIES = Sheet(
    title="Компании",
    file_prefix="companies",
    columns=COMPANY_COLUMNS,
    row=company_row,
    prepare=_prepare_companies,
    widths={
        "Компания": 35,
        "LinkedIn slug": 30,
        "Ссылка на компанию": 45,
        "Сайт": 35,
        "Отрасль": 30,
        "Размер": 26,
        "Штаб-квартира": 30,
        "Ошибка загрузки профиля": 40,
        "Логотип": 40,
        "Короткое описание": 50,
        "Описание компании": 80,
    },
)


def _filename(sheet: Sheet, extension: str) -> str:
    return f"{sheet.file_prefix}_{timezone.localtime().strftime('%Y%m%d_%H%M')}.{extension}"


def queryset_for_export(queryset, sheet: Sheet = VACANCIES):
    return sheet.prepare(queryset)


def export_csv(queryset, sheet: Sheet = VACANCIES) -> HttpResponse:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow([label for _key, label in sheet.columns])

    for obj in sheet.prepare(queryset).iterator(chunk_size=200):
        row = sheet.row(obj)
        writer.writerow(["" if row[key] is None else str(row[key]) for key, _label in sheet.columns])

    # utf-8-sig - чтобы Excel не ломал кириллицу при двойном клике
    response = HttpResponse(buffer.getvalue().encode("utf-8-sig"), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(sheet, "csv")}"'
    return response


def export_xlsx(queryset, sheet: Sheet = VACANCIES) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=False)
    worksheet = workbook.active
    worksheet.title = sheet.title

    worksheet.append([label for _key, label in sheet.columns])
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"

    for obj in sheet.prepare(queryset).iterator(chunk_size=200):
        row = sheet.row(obj)
        values = []
        for key, _label in sheet.columns:
            value = row[key]
            if value is None:
                value = ""
            elif not isinstance(value, (int, float, datetime)) and not isinstance(value, str):
                value = str(value)
            if isinstance(value, str):
                # в ячейку Excel влезает 32767 символов
                value = value[:32000]
            values.append(value)
        worksheet.append(values)

    for index, (_key, label) in enumerate(sheet.columns, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = sheet.widths.get(label, 20)

    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(sheet, "xlsx")}"'
    return response
